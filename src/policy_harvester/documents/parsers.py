from __future__ import annotations

import hashlib
import io
import re
import subprocess
import tempfile
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Protocol

import olefile
import pymupdf as fitz
from docx import Document as DocxDocument
from lxml import etree, html
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from .types import DetectedType, detect_type


@dataclass(frozen=True)
class Block:
    stable_key: str
    kind: str
    text: str = ""
    page_number: int | None = None
    source_path: str | None = None
    bbox: tuple[float, float, float, float] | None = None
    table_data: dict | None = None
    ocr_used: bool = False
    ocr_confidence: float | None = None
    metadata: dict = field(default_factory=dict)


@dataclass(frozen=True)
class ParseResult:
    parser_name: str
    parser_version: str
    status: str
    blocks: tuple[Block, ...]
    text: str
    quality_flags: tuple[str, ...] = ()
    # Pictures to be transcribed by the multimodal LLM. Each one has a placeholder block
    # (kind "image", metadata["pending_image"] == key) marking where its text belongs.
    images: tuple["EmbeddedImage", ...] = ()
    # Files derived from the source for display, e.g. the PDF rendering of an HWP document.
    derived: tuple["DerivedFile", ...] = ()


@dataclass(frozen=True)
class DerivedFile:
    kind: str  # "pdf_render"
    data: bytes
    tool: str  # "rhwp 0.8.7"


@dataclass(frozen=True)
class EmbeddedImage:
    key: str
    data: bytes
    page_number: int | None = None
    metadata: dict = field(default_factory=dict)


def image_placeholder(stable_key: str, image_key: str, source_path: str,
                      page_number: int | None = None, metadata: dict | None = None) -> Block:
    return Block(stable_key, "image", "", page_number=page_number, source_path=source_path,
                 metadata={**(metadata or {}), "pending_image": image_key})


class DocumentParser(Protocol):
    formats: frozenset[str]
    name: str
    version: str
    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult: ...


def _key(prefix: str, index: int) -> str:
    return f"{prefix}:{index:06d}"


BLOCK_TAGS = frozenset({"p", "div", "li", "td", "th", "tr", "table", "h1", "h2", "h3", "h4",
                        "h5", "h6", "ul", "ol", "dl", "dt", "dd", "section", "article", "blockquote"})


def _mark_text_boundaries(root) -> None:
    """Separate block elements and <br>, but never inline runs: "1<span>7:00</span>" is 17:00."""
    for node in root.iter():
        if not isinstance(node.tag, str):
            continue
        if node.tag == "br":
            node.tail = "\n" + (node.tail or "")
        elif node.tag in BLOCK_TAGS:
            node.tail = " " + (node.tail or "")


def _cell_text(node) -> str:
    return re.sub(r"\s+", " ", "".join(node.itertext())).strip()


def _own_rows(table) -> list:
    # Rows of nested tables belong to those tables, not to this one.
    return [row for row in table.xpath(".//tr") if next(row.iterancestors("table")) is table]


def _span(value: str | None) -> int:
    value = (value or "").strip()
    return max(1, int(value)) if value.isdigit() else 1


def _is_layout_table(table) -> bool:
    return all(len(row.xpath("./th|./td")) <= 1 for row in _own_rows(table))


def table_grid(table, *, mark_spans: bool = False) -> dict:
    """Expand rowspan/colspan into a rectangular grid and keep the original cell spans.

    With ``mark_spans`` the text writes a merged cell once, on its first row, tagged with the
    rows/columns it covers (e.g. "34,000 [4행 병합]"), so a shared value is not read as one value
    per row or per column; rows it also covers show "↑" in its place.
    """
    grid: dict[tuple[int, int], int] = {}
    cells = []
    for row_index, row in enumerate(_own_rows(table)):
        column = 0
        for cell in row.xpath("./th|./td"):
            while (row_index, column) in grid:
                column += 1
            row_span, col_span = _span(cell.get("rowspan")), _span(cell.get("colspan"))
            cells.append({"row": row_index, "col": column, "rowspan": row_span,
                          "colspan": col_span, "header": cell.tag == "th", "text": _cell_text(cell)})
            for dy in range(min(row_span, 1000)):
                for dx in range(min(col_span, 100)):
                    grid[(row_index + dy, column + dx)] = len(cells) - 1
            column += col_span
    height = max((key[0] for key in grid), default=-1) + 1
    width = max((key[1] for key in grid), default=-1) + 1
    origins = [[grid.get((y, x)) for x in range(width)] for y in range(height)]
    rows = [[cells[origin]["text"] if origin is not None else "" for origin in row] for row in origins]
    # Text keeps one value per source cell in a row, so colspans do not repeat but equal
    # values from different cells (e.g. "O | O | X") survive. Empty cells keep their place so a
    # later value is not read as belonging to an earlier column.
    lines = []
    for row_index, row in enumerate(origins):
        seen: list[int] = []
        values: list[str] = []
        for origin in row:
            if origin is None:
                values.append("")
            elif origin not in seen:
                seen.append(origin)
                cell = cells[origin]
                if not mark_spans:
                    values.append(cell["text"])
                elif cell["row"] != row_index:
                    values.append("↑" if cell["text"] else "")
                else:
                    values.append(_span_text(cell))
        while values and not values[-1]:
            values.pop()
        if any(value and value != "↑" for value in values):
            lines.append(" | ".join(values))
    return {"rows": rows, "cells": cells, "text": "\n".join(lines)}


def _span_text(cell: dict) -> str:
    spans = [f"{cell['rowspan']}행" if cell["rowspan"] > 1 else "",
             f"{cell['colspan']}열" if cell["colspan"] > 1 else ""]
    label = "·".join(part for part in spans if part)
    return f"{cell['text']} [{label} 병합]" if label else cell["text"]


STRUCK = "//del|//s|//strike|//*[contains(translate(@style, ' ', ''), 'line-through')]"


def _mark_strikethrough(root) -> None:
    """Wrap struck-out text in ~~ ~~ so an old deadline is not read as a current one."""
    struck = root.xpath(STRUCK)
    inside = {id(node) for node in struck}
    for node in struck:
        if any(id(parent) in inside for parent in node.iterancestors()):
            continue
        if not "".join(node.itertext()).strip():
            continue
        node.text = "~~" + (node.text or "")
        if len(node):
            node[-1].tail = (node[-1].tail or "") + "~~"
        else:
            node.text += "~~"


class HtmlParser:
    formats = frozenset({"html"})
    name, version = "lxml-html", "1.4"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        try:
            payload.decode("utf-8")
            # Without a <meta charset>, lxml would otherwise decode bytes as Latin-1.
            root = html.fromstring(payload, parser=html.HTMLParser(encoding="utf-8", huge_tree=True))
        except UnicodeDecodeError:
            root = html.fromstring(payload, parser=html.HTMLParser(huge_tree=True))
        _mark_strikethrough(root)
        _mark_text_boundaries(root)
        blocks: list[Block] = []
        xpath = root.getroottree().getpath
        kinds = {"h1": "heading", "h2": "heading", "h3": "heading", "h4": "heading",
                 "p": "paragraph", "li": "list_item", "table": "table"}
        for node in root.xpath("//h1|//h2|//h3|//h4|//p|//li|//table"):
            # A data table block already carries its cell text; layout tables do not.
            if any(not _is_layout_table(table) for table in node.iterancestors("table")):
                continue
            table_data = None
            if node.tag == "table":
                if _is_layout_table(node):
                    continue
                table_data = table_grid(node)
                text_value = table_data.pop("text")
            else:
                text_value = _cell_text(node)
            if not text_value:
                continue
            blocks.append(Block(_key("html", len(blocks)), kinds[node.tag], text_value,
                                source_path=xpath(node), table_data=table_data))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


def normalized_bbox(x0: float, y0: float, x1: float, y1: float
                    ) -> tuple[float, float, float, float] | None:
    """Clamp to the page; text drawn partly off-page or with zero area has no usable box."""
    box = tuple(min(1.0, max(0.0, float(value))) for value in (x0, y0, x1, y1))
    return box if box[0] < box[2] and box[1] < box[3] else None


class PdfParser:
    """Text layer via PyMuPDF; pages without a text layer are rendered for LLM transcription."""
    formats = frozenset({"pdf"})
    name, version = "pymupdf-vision", "2.0"
    min_text_chars = 30

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        blocks: list[Block] = []
        images: list[EmbeddedImage] = []
        with fitz.open(stream=payload, filetype="pdf") as document:
            for page_index, page in enumerate(document):
                page_blocks = page.get_text("blocks", sort=True)
                if sum(len(str(item[4]).strip()) for item in page_blocks) < self.min_text_chars:
                    key = f"page-{page_index + 1}"
                    png = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
                    images.append(EmbeddedImage(key, png, page_index + 1, {"rendered_page": page_index + 1}))
                    blocks.append(image_placeholder(_key("pdf", len(blocks)), key,
                                                    f"page/{page_index + 1}/image", page_index + 1))
                    continue
                width, height = page.rect.width, page.rect.height
                for raw in page_blocks:
                    text_value = re.sub(r"\s+", " ", str(raw[4])).strip()
                    if not text_value:
                        continue
                    bbox = normalized_bbox(raw[0] / width, raw[1] / height,
                                           raw[2] / width, raw[3] / height)
                    blocks.append(Block(_key("pdf", len(blocks)), "paragraph", text_value,
                                        page_number=page_index + 1,
                                        source_path=f"page/{page_index + 1}/block/{len(blocks)}",
                                        bbox=bbox,
                                        metadata={} if bbox else {"raw_bbox_points": list(raw[:4])}))
        text_value = "\n\n".join(block.text for block in blocks if block.text)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), text_value,
                           ("scanned_pages",) if images else (), tuple(images))


class ImageParser:
    """Posters and photos are transcribed whole by the multimodal LLM (see worker)."""
    formats = frozenset({"png", "jpeg", "gif"})
    name, version = "llm-vision", "2.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        return ParseResult(self.name, self.version, "succeeded",
                           (image_placeholder(_key("image", 0), "image", "image"),), "", (),
                           (EmbeddedImage("image", payload),))


class DocxParser:
    formats = frozenset({"docx"})
    name, version = "python-docx", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        doc = DocxDocument(io.BytesIO(payload))
        blocks: list[Block] = []
        for paragraph_index, paragraph in enumerate(doc.paragraphs):
            text_value = paragraph.text.strip()
            if text_value:
                kind = "heading" if paragraph.style and paragraph.style.name.startswith("Heading") else "paragraph"
                blocks.append(Block(_key("docx", len(blocks)), kind, text_value,
                                    source_path=f"paragraph/{paragraph_index}"))
        for table_index, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text_value = "\n".join("\t".join(row) for row in rows)
            blocks.append(Block(_key("docx", len(blocks)), "table", text_value,
                                source_path=f"table/{table_index}", table_data={"rows": rows}))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


class HwpxParser:
    formats = frozenset({"hwpx"})
    name, version = "hwpx-xml", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        blocks: list[Block] = []
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            section_names = sorted(name for name in archive.namelist()
                                   if name.startswith("Contents/section") and name.endswith(".xml"))
            parser = etree.XMLParser(resolve_entities=False, no_network=True, recover=True)
            for section_index, name in enumerate(section_names):
                root = etree.fromstring(archive.read(name), parser)
                paragraphs = root.xpath("//*[local-name()='p']")
                for paragraph_index, paragraph in enumerate(paragraphs):
                    text_value = "".join(paragraph.xpath(".//*[local-name()='t']/text()" )).strip()
                    if text_value:
                        blocks.append(Block(_key("hwpx", len(blocks)), "paragraph", text_value,
                                            source_path=f"{name}/p/{paragraph_index}"))
                tables = root.xpath("//*[local-name()='tbl']")
                for table_index, table in enumerate(tables):
                    rows = []
                    for row in table.xpath("./*[local-name()='tr']"):
                        rows.append(["".join(cell.xpath(".//*[local-name()='t']/text()")).strip()
                                     for cell in row.xpath("./*[local-name()='tc']")])
                    text_value = "\n".join("\t".join(row) for row in rows)
                    blocks.append(Block(_key("hwpx", len(blocks)), "table", text_value,
                                        source_path=f"{name}/table/{table_index}", table_data={"rows": rows}))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


def _compact(value: str) -> str:
    # Symbol-font bullets live in the private-use area and carry no text.
    return re.sub(r"[\s\ue000-\uf8ff]+", "", value)


class HwpParser:
    """hwp5html for structure and tables, hwp5txt to recover dropped lines, BinData OCR.

    hwp5html silently drops some paragraphs (for example ones holding a hyperlink field), and
    neither converter emits embedded pictures, so both gaps are filled from other sources and
    every added block records where it came from.
    """
    formats = frozenset({"hwp"})
    name, version = "pyhwp-hwp5html", "2.3"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        with tempfile.TemporaryDirectory(prefix="policy-hwp-") as directory:
            source = Path(directory) / "document.hwp"
            target = Path(directory) / "document.html"
            source.write_bytes(payload)
            flags: list[str] = []
            try:  # large documents can take minutes; plain text below still covers them
                html_run = subprocess.run(["hwp5html", "--html", "--output", str(target), str(source)],
                                          capture_output=True, check=False, timeout=300)
                html_bytes = target.read_bytes() if html_run.returncode == 0 and target.exists() else None
            except subprocess.TimeoutExpired:
                html_bytes = None
                flags.append("hwp5html_timeout")
            text_run = subprocess.run(["hwp5txt", str(source)], capture_output=True,
                                      check=False, timeout=180)
        plain = text_run.stdout.decode("utf-8", "replace") if text_run.returncode == 0 else ""
        blocks: list[Block] = []
        if html_bytes is not None:
            converted = HtmlParser().parse(html_bytes)
            blocks.extend(Block("", block.kind, block.text, source_path=f"hwp5html{block.source_path}",
                                table_data=block.table_data) for block in converted.blocks)
            seen = _compact(converted.text)
            for index, line in enumerate(plain.splitlines()):
                key = _compact(line)
                if len(key) < 4 or key in {"<표>", "<그림>"} or key in seen:
                    continue
                seen += key
                blocks.append(Block("", "paragraph", line.strip(), source_path=f"hwp5txt/line/{index}",
                                    metadata={"recovered_from": "hwp5txt",
                                              "reason": "absent from hwp5html output"}))
                if "hwp5txt_recovered_lines" not in flags:
                    flags.append("hwp5txt_recovered_lines")
            status = "succeeded"
        else:
            if text_run.returncode:
                raise ValueError(text_run.stderr.decode("utf-8", "replace") or "hwp5html and hwp5txt failed")
            reason = html_run.stderr.decode("utf-8", "replace")[-300:]
            blocks.extend(Block("", "paragraph", value.strip(), source_path=f"paragraph/{index}",
                                metadata={"hwp5html_error": reason})
                          for index, value in enumerate(re.split(r"\n\s*\n", plain)) if value.strip())
            # hwp5txt replaces tables with a placeholder, so the result is only partial.
            status = "partial"
            flags.append("hwp_tables_lost")
        images = self._embedded_images(payload)
        # Pictures have no anchor in the converted text, so they follow the body.
        blocks.extend(image_placeholder("", image.key, f"bindata/{image.key}", metadata=image.metadata)
                      for image in images)
        final = tuple(Block(_key("hwp", index), block.kind, block.text,
                            source_path=block.source_path, table_data=block.table_data,
                            ocr_used=block.ocr_used, ocr_confidence=block.ocr_confidence,
                            metadata=block.metadata) for index, block in enumerate(blocks))
        return ParseResult(self.name, self.version, status, final,
                           "\n\n".join(block.text for block in final if block.text),
                           tuple(sorted(set(flags))), tuple(images))

    def _embedded_images(self, payload: bytes) -> list[EmbeddedImage]:
        """Pictures stored in the OLE BinData storage (zlib raw deflate when compressed)."""
        try:
            document = olefile.OleFileIO(io.BytesIO(payload))
        except (OSError, IOError):
            return []
        images = []
        with document:
            header = document.openstream("FileHeader").read() if document.exists("FileHeader") else b""
            compressed = len(header) >= 40 and bool(int.from_bytes(header[36:40], "little") & 1)
            for path in document.listdir(streams=True):
                if len(path) != 2 or path[0] != "BinData":
                    continue
                data = document.openstream(path).read()
                if compressed:
                    try:
                        data = zlib.decompress(data, -15)
                    except zlib.error:
                        pass
                try:
                    with Image.open(io.BytesIO(data)) as image:
                        image.verify()
                except Exception:  # noqa: BLE001 - OLE objects, fonts and other non-pictures
                    continue
                images.append(EmbeddedImage(path[1], data, metadata={"hwp_bindata": path[1]}))
        return images


class DocParser:
    """Legacy Word 97-2003 via antiword; tables come out as text rows, not structure."""
    formats = frozenset({"doc"})
    name, version = "antiword", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        with tempfile.TemporaryDirectory(prefix="policy-doc-") as directory:
            source = Path(directory) / "document.doc"
            source.write_bytes(payload)
            process = subprocess.run(["antiword", "-w", "0", "-m", "UTF-8.txt", str(source)],
                                     capture_output=True, check=False, timeout=60)
        if process.returncode:
            raise ValueError(process.stderr.decode("utf-8", "replace")[-300:] or "antiword failed")
        text_value = process.stdout.decode("utf-8", "replace")
        blocks = tuple(Block(_key("doc", index), "paragraph", re.sub(r"[ \t]+", " ", value).strip(),
                             source_path=f"paragraph/{index}")
                       for index, value in enumerate(re.split(r"\n\s*\n", text_value)) if value.strip())
        return ParseResult(self.name, self.version, "succeeded", blocks,
                           "\n\n".join(block.text for block in blocks), ("doc_tables_as_text",))


class XlsxParser:
    formats = frozenset({"xlsx"})
    name, version = "openpyxl", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        blocks = []
        for worksheet in workbook.worksheets:
            rows = [["" if value is None else str(value) for value in row]
                    for row in worksheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(value for value in row)]
            if rows:
                blocks.append(Block(_key("xlsx", len(blocks)), "table",
                                    "\n".join("\t".join(row) for row in rows),
                                    source_path=f"sheet/{worksheet.title}", table_data={"rows": rows}))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


class XlsParser:
    """Legacy Excel 97-2003 (often blank application forms) via xlrd, one table per sheet."""
    formats = frozenset({"xls"})
    name, version = "xlrd", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=payload)
        blocks = []
        for sheet in workbook.sheets():
            rows = [[_xls_cell(sheet.cell(row, column), workbook.datemode) for column in range(sheet.ncols)]
                    for row in range(sheet.nrows)]
            rows = [row for row in rows if any(value for value in row)]
            if rows:
                blocks.append(Block(_key("xls", len(blocks)), "table",
                                    "\n".join("\t".join(row) for row in rows),
                                    source_path=f"sheet/{sheet.name}", table_data={"rows": rows}))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


def _xls_cell(cell: Any, datemode: int) -> str:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode).isoformat(sep=" ").removesuffix(" 00:00:00")
        except (ValueError, OverflowError):
            return str(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return str(int(cell.value))
    return "" if cell.value is None else str(cell.value).strip()


class PptxParser:
    formats = frozenset({"pptx"})
    name, version = "python-pptx", "1.0"

    def parse(self, payload: bytes, filename: str | None = None) -> ParseResult:
        presentation = Presentation(io.BytesIO(payload))
        blocks = []
        for slide_index, slide in enumerate(presentation.slides):
            for shape_index, shape in enumerate(slide.shapes):
                text_value = getattr(shape, "text", "").strip()
                if text_value:
                    blocks.append(Block(_key("pptx", len(blocks)), "paragraph", text_value,
                                        page_number=slide_index + 1,
                                        source_path=f"slide/{slide_index + 1}/shape/{shape_index}"))
        joined = "\n\n".join(block.text for block in blocks)
        return ParseResult(self.name, self.version, "succeeded", tuple(blocks), joined)


class ParserRegistry:
    def __init__(self):
        parsers: tuple[DocumentParser, ...] = (
            HtmlParser(), PdfParser(), ImageParser(), DocxParser(), DocParser(), HwpxParser(), HwpParser(),
            XlsxParser(), XlsParser(), PptxParser(),
        )
        self._parsers = {format_name: parser for parser in parsers for format_name in parser.formats}
        self.pdf = PdfHybridParser()
        self.hwp = HwpPdfParser({"hwp": self._parsers["hwp"], "hwpx": self._parsers["hwpx"]})

    def parse(self, payload: bytes, filename: str | None = None,
              detected: DetectedType | None = None) -> ParseResult:
        kind = detected or detect_type(payload, filename)
        if kind.format == "zip":
            return self._parse_archive(payload)
        if kind.format == "pdf":
            return self._with_fallback(lambda: self.pdf.parse(payload, filename), kind.format,
                                       payload, filename, "pdf_hybrid_failed")
        if kind.format in HwpPdfParser.formats:
            return self._with_fallback(lambda: self.hwp.parse(payload, filename, format_name=kind.format),
                                       kind.format, payload, filename, "rhwp_failed")
        if kind.format in DOCLING_ONLY:
            try:
                return DoclingParser().parse(payload, filename, format_name=kind.format)
            except Exception as exc:  # noqa: BLE001 - damaged files fall back to the simple parser
                fallback = self._parsers[kind.format].parse(payload, filename)
                return ParseResult(fallback.parser_name, fallback.parser_version, "partial",
                                   fallback.blocks, fallback.text,
                                   (*fallback.quality_flags, f"docling_failed:{exc.__class__.__name__}"),
                                   fallback.images)
        parser = self._parsers.get(kind.format)
        if parser is None:
            return ParseResult("unsupported", "1.0", "unsupported", (), "", (f"unsupported:{kind.format}",))
        return parser.parse(payload, filename)

    def _with_fallback(self, parse, format_name: str, payload: bytes, filename: str | None,
                       flag: str) -> ParseResult:
        """The preferred route; on failure the direct parser, reported as partial."""
        try:
            return parse()
        except Exception as exc:  # noqa: BLE001 - damaged files still get the simple parser
            fallback = self._parsers[format_name].parse(payload, filename)
            return ParseResult(fallback.parser_name, fallback.parser_version, "partial",
                               fallback.blocks, fallback.text,
                               (*fallback.quality_flags, f"{flag}:{exc.__class__.__name__}"),
                               fallback.images)

    def identity(self, detected: DetectedType) -> tuple[str, str]:
        if detected.format == "zip":
            return "safe-zip-recursive", "1.1"
        if detected.format == "pdf":
            return self.pdf.name, self.pdf.version
        if detected.format in HwpPdfParser.formats:
            return self.hwp.name, self.hwp.version
        if detected.format in DOCLING_ONLY:
            return DoclingParser.name, DoclingParser.version
        parser = self._parsers.get(detected.format)
        if parser is None:
            return "unsupported", "1.0"
        return parser.name, parser.version

    def _parse_archive(self, payload: bytes) -> ParseResult:
        blocks: list[Block] = []
        flags: list[str] = []
        images: list[EmbeddedImage] = []
        total_uncompressed = 0
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = [item for item in archive.infolist() if not item.is_dir()]
            if len(members) > 250:
                raise ValueError("archive member limit exceeded")
            for member in members:
                total_uncompressed += member.file_size
                if total_uncompressed > 200 * 1024 * 1024:
                    raise ValueError("archive expanded size limit exceeded")
                if member.file_size > 50 * 1024 * 1024:
                    flags.append(f"member_too_large:{member.filename}")
                    continue
                try:
                    member_payload = archive.read(member)
                except (RuntimeError, OSError, zipfile.BadZipFile) as exc:
                    flags.append(f"unreadable_member:{member.filename}:{exc.__class__.__name__}")
                    continue
                detected = detect_type(member_payload, member.filename)
                if detected.format in {"zip", "binary"}:
                    flags.append(f"unsupported_member:{member.filename}")
                    continue
                result = self.parse(member_payload, member.filename, detected)
                prefix = f"{member.filename}#"
                images.extend(EmbeddedImage(prefix + image.key, image.data, image.page_number,
                                            {**image.metadata, "archive_member": member.filename})
                              for image in result.images)
                for block in result.blocks:
                    blocks.append(Block(
                        stable_key=_key("archive", len(blocks)),
                        kind=block.kind,
                        text=block.text,
                        page_number=block.page_number,
                        source_path=f"member/{member.filename}/{block.source_path or block.stable_key}",
                        bbox=block.bbox,
                        table_data=block.table_data,
                        ocr_used=block.ocr_used,
                        ocr_confidence=block.ocr_confidence,
                        metadata={**block.metadata, "archive_member": member.filename,
                                  "member_parser": result.parser_name,
                                  **({"pending_image": prefix + block.metadata["pending_image"]}
                                     if "pending_image" in block.metadata else {})},
                    ))
                flags.extend(result.quality_flags)
        joined = "\n\n".join(block.text for block in blocks)
        status = "partial" if flags else "succeeded"
        return ParseResult("safe-zip-recursive", "1.1", status, tuple(blocks), joined,
                           tuple(sorted(set(flags))), tuple(images))

    @property
    def supported_formats(self) -> frozenset[str]:
        return frozenset(self._parsers)


from .docling_parser import DoclingParser  # noqa: E402 - imports the shared types above
from .hwp_pdf import HwpPdfParser  # noqa: E402
from .pdf_hybrid import PdfHybridParser  # noqa: E402

DOCLING_ONLY = frozenset({"docx", "pptx", "xlsx"})
